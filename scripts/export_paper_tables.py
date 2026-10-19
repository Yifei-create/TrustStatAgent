from __future__ import annotations

import json
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill

ROOT = Path(__file__).resolve().parents[1]
NUMERIC = ROOT / "results/metrics/numeric_metrics.jsonl"
STATQA = ROOT / "results/metrics/statqa_mini_ranker_opt_full.jsonl"
STATQA_BREAKDOWN = ROOT / "results/metrics/statqa_category_breakdown.json"
OUTPUT = ROOT / "results/tables/TrustStatAgent_ESWA_tables.xlsx"


def read_jsonl(path: Path) -> list[dict]:
    return [json.loads(line) for line in path.read_text(encoding="utf-8").splitlines() if line.strip()]


def mean_value(rows: list[dict], key: str):
    values = [float(row[key]) for row in rows if row.get(key) is not None]
    return sum(values) / len(values) if values else "-"


def pct(rows: list[dict], key: str):
    value = mean_value(rows, key)
    return "-" if value == "-" else 100.0 * value


def pct_optional(row: dict, key: str):
    value = row.get(key)
    return "-" if value is None else 100.0 * float(value)


def add_sheet(workbook: Workbook, name: str, headers: list[str], rows: list[list[object]]) -> None:
    sheet = workbook.create_sheet(name)
    sheet.append(headers)
    for row in rows:
        sheet.append(row)
    fill = PatternFill("solid", fgColor="D9EAF7")
    for cell in sheet[1]:
        cell.font = Font(bold=True)
        cell.fill = fill
        cell.alignment = Alignment(horizontal="center")
    for column in sheet.columns:
        sheet.column_dimensions[column[0].column_letter].width = min(max(len(str(cell.value or "")) for cell in column) + 2, 36)
    sheet.freeze_panes = "A2"


def main() -> None:
    numeric = read_jsonl(NUMERIC)
    statqa = read_jsonl(STATQA)
    category_breakdown = json.loads(STATQA_BREAKDOWN.read_text(encoding="utf-8"))
    ng: dict[tuple[str, str], list[dict]] = {}
    for row in numeric:
        ng.setdefault((row["dataset"], row["variant"]), []).append(row)
    sg = {row["variant"]: row for row in statqa}
    trust_statqa = sg["TrustStatAgent"]
    react_statqa = sg["ReAct with Conventional Gating"]
    noplan_rows = read_jsonl(ROOT / "results/raw/statqa_mini_without_explicit_planning_opt_mini_full_test.jsonl")
    noplan_accuracy = sum(bool(row["official_correct"]) for row in noplan_rows) / len(noplan_rows)

    trust = "TrustStatAgent"
    conventional = "ReAct with Conventional Gating"
    cali, carbon = "California Housing", "Carbon Monitor China/Total"
    workbook = Workbook()
    workbook.remove(workbook.active)

    add_sheet(workbook, "Table1_Task", ["Dataset", "Task", "Metric", "TrustStatAgent"], [
        ["StatQA", "Statistical QA", "Official Accuracy (%)", 100 * trust_statqa["accuracy"]],
        ["Adult", "Binary classification", "Accuracy (%)", pct(ng.get(("Adult", trust), []), "accuracy")],
        ["California Housing", "Regression", "MAE", mean_value(ng.get((cali, trust), []), "mae")],
        ["Carbon Monitor", "One-day forecasting", "MAE", mean_value(ng.get((carbon, trust), []), "mae")],
    ])
    add_sheet(workbook, "Table1_StatQA_Baselines", ["Method", "Official Accuracy (%)"], [
        ["Zero-shot-CoT", "-"], ["ReAct Statistical Agent", "-"], ["ReAct + Conventional Gate", "-"], ["TrustStatAgent", 100 * trust_statqa["accuracy"]],
    ])

    table2 = [["StatQA", "ReAct Statistical Agent", "-", "-", 100.0, 0.0, 100 * (1 - react_statqa["accuracy"])]]
    for label, dataset in [("Adult", "Adult"), ("California", cali), ("Carbon", carbon)]:
        direct_rows = ng.get((dataset, "w/o Risk Gate"), [])
        table2.append([label, "ReAct Statistical Agent", "-", "-", 100.0, 0.0, pct(direct_rows, "error_release_rate")])
    for label, dataset in [("StatQA", "StatQA"), ("Adult", "Adult"), ("California", cali), ("Carbon", carbon)]:
        conv = react_statqa if dataset == "StatQA" else ng.get((dataset, conventional), [{}])[0]
        full = trust_statqa if dataset == "StatQA" else ng.get((dataset, trust), [{}])[0]
        table2.extend([
            [label, conventional, "-" if dataset == "StatQA" else pct_optional(conv, "coverage"), "-" if dataset == "StatQA" else pct_optional(conv, "released_correctness"), 100 * conv["release_rate"], 100 * conv["manual_review_rate"], 100 * conv["error_release_rate"]],
            [label, trust, "-" if dataset == "StatQA" else pct_optional(full, "coverage"), "-" if dataset == "StatQA" else pct_optional(full, "released_correctness"), 100 * full["release_rate"], 100 * full["manual_review_rate"], 100 * full["error_release_rate"]],
        ])
    add_sheet(workbook, "Table2_Baseline_Reliability", ["Dataset", "Method", "Conformal coverage (%)", "Released correctness (%)", "Auto-release (%)", "Manual-review (%)", "Error-release (%)"], table2)

    add_sheet(workbook, "Table3_Planning_Ablation", ["Variant", "Official Accuracy (%)"], [
        ["w/o Explicit Planning", 100 * noplan_accuracy], ["Full TrustStatAgent", 100 * trust_statqa["accuracy"]],
    ])

    ablation = []
    for label, dataset in [("StatQA", "StatQA"), ("Adult", "Adult"), ("California", cali), ("Carbon", carbon)]:
        if dataset == "StatQA":
            rows = {row["variant"]: row for row in statqa}
            variants = [rows["w/o Conformal Calibration"], rows["w/o Risk Gate"], rows["TrustStatAgent"]]
        else:
            variants = [ng.get((dataset, "w/o Conformal Calibration"), [{}])[0], ng.get((dataset, "w/o Risk Gate"), [{}])[0], ng.get((dataset, trust), [{}])[0]]
        for row in variants:
            ablation.append([label, row["variant"], "-" if dataset == "StatQA" else pct_optional(row, "coverage"), "-" if dataset == "StatQA" else pct_optional(row, "released_correctness"), 100 * row["release_rate"], 100 * row["manual_review_rate"], 100 * row["error_release_rate"]])
    add_sheet(workbook, "Table4_Module_Ablation", ["Dataset", "Variant", "Conformal coverage (%)", "Released correctness (%)", "Auto-release (%)", "Manual-review (%)", "Error-release (%)"], ablation)

    category_rows = []
    for row in category_breakdown["rows"]:
        category_rows.append([
            row["category"], row["n"], 100 * row["accuracy"], 100 * row["release_rate"],
            100 * row["manual_review_rate"], 100 * row["error_release_rate"], 100 * row["balanced_accuracy"],
        ])
    add_sheet(workbook, "Table5_StatQA_Categories", ["Category", "N", "Accuracy (%)", "Auto-release (%)", "Manual-review (%)", "Error-release (%)", "Balanced Accuracy (%)"], category_rows)

    notes = workbook.create_sheet("Notes")
    notes.append(["Status"])
    for note in [
        "StatQA mini uses a fixed TF-IDF plus logistic ranker trained on scorer-calibration labels only.",
        "StatQA conformal calibration uses a disjoint split; alpha is selected once on validation and test is held out.",
        "StatQA baseline task accuracy rows are intentionally omitted from the main task table; reliability baselines are shown in Table2.",
        "Numeric dataset values are the existing formal results and were not rerun in this StatQA experiment.",
        "Error-release is the error rate among automatically released results.",
        "Conformal coverage is unconditional prediction-set or interval coverage; released correctness is conditional on automatic release.",
    ]:
        notes.append([note])
    notes.column_dimensions["A"].width = 120
    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    main()
